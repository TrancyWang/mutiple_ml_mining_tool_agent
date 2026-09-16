"""文本挖掘与机器学习工具集 - Qwen-Agent Function Calling Tools

将客户端的所有功能封装为可调用的工具函数，供 Agent 通过 Function Calling 调用。
"""

import os
import sys
import json
from typing import Optional, Dict, List, Any
from pathlib import Path
import time
import ast
from academic_agent.infrastructure.runtime_safety import configure_numeric_runtime

configure_numeric_runtime()

import pandas as pd
import numpy as np
from academic_agent.infrastructure.workspace_manager import workspace_manager
from academic_agent.infrastructure.analysis_artifacts import (
    artifact_output_root,
    artifact_result,
    create_run_dir,
    infer_source_scope,
    write_workbook,
)

# 本项目算法层统一放在 algorithms，避免与源项目 video_text_mutiplemodal_agent
# 的 src_codes 目录混淆。
BASE_DIR = Path(__file__).parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from academic_agent.integrations.video_text_adapter import source_video_text_adapter


class TextMiningTools:
    """文本挖掘工具集"""
    
    def __init__(self):
        """初始化工具集"""
        # 当前工作数据
        self.current_data = None
        self.current_file_path = None
        self.current_source_scope = "upload"
        self.engine_name = "video_text_mutiplemodal_agent"

    def _get_output_path(self, original_name: str, suffix: str) -> str:
        """生成带时间戳的唯一输出文件路径"""
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        root = artifact_output_root(self.current_file_path, self.current_source_scope)
        return str(root / f"{Path(original_name).stem}_{suffix}_{timestamp}.csv")

    @staticmethod
    def _runtime_log(tool: str, message: str) -> None:
        from academic_agent.agent.executor import emit_tool_log

        emit_tool_log(tool, message)

    def _detect_id_column(self) -> Optional[str]:
        if self.current_data is None:
            return None
        for candidate in ("cid", "id", "ID", "编号", "序号"):
            if candidate in self.current_data.columns:
                return candidate
        return None

    def auto_detect_text_column(self) -> Optional[str]:
        """自动检测可能包含文本的列"""
        if self.current_data is None or self.current_data.empty:
            return None
            
        text_column_candidates = ['text', 'content', 'comment', 'review', '文本', '内容', '评论']
        for col in self.current_data.columns:
            if str(col).lower() in text_column_candidates:
                return str(col)
                
        longest_col = None
        max_avg_len = 0
        for col in self.current_data.columns:
            if self.current_data[col].dtype == 'object':
                avg_len = self.current_data[col].dropna().astype(str).map(len).mean()
                if avg_len > max_avg_len:
                    max_avg_len = avg_len
                    longest_col = col
                    
        return longest_col

    def get_data_info(self) -> Dict[str, Any]:
        """Return a lightweight schema summary for the registered data tool."""
        if self.current_data is None:
            return {"success": False, "error": "请先加载数据"}
        return {
            "success": True,
            "rows": int(len(self.current_data)),
            "columns": int(len(self.current_data.columns)),
            "column_names": [str(column) for column in self.current_data.columns],
            "dtypes": {str(column): str(dtype) for column, dtype in self.current_data.dtypes.items()},
            "text_column": self.auto_detect_text_column(),
        }

    def load_data(
        self, file_path: str, encoding: str = "UTF-8", source_scope: str | None = None
    ) -> Dict[str, Any]:
        """按扩展名加载常见学术数据、文本和文档文件。"""
        try:
            if not os.path.exists(file_path):
                return {"success": False, "error": f"文件不存在: {file_path}"}

            suffix = Path(file_path).suffix.lower()
            if suffix in {'.csv', '.tsv'}:
                separator = '\t' if suffix == '.tsv' else None
                try:
                    if separator:
                        df = pd.read_csv(file_path, encoding=encoding, sep=separator)
                    else:
                        df = pd.read_csv(file_path, encoding=encoding)
                except UnicodeDecodeError:
                    if separator:
                        df = pd.read_csv(file_path, encoding='gb18030', sep=separator)
                    else:
                        df = pd.read_csv(file_path, encoding='gb18030')
            elif suffix in {'.xlsx', '.xls', '.xlsm', '.ods'}:
                df = pd.read_excel(file_path)
            elif suffix in {'.txt', '.log', '.md', '.rst'}:
                text = Path(file_path).read_text(encoding=encoding, errors='replace')
                df = pd.DataFrame({'text': [line for line in text.splitlines() if line.strip()]})
            elif suffix in {'.json', '.jsonl', '.ndjson'}:
                if suffix in {'.jsonl', '.ndjson'}:
                    df = pd.read_json(file_path, lines=True)
                else:
                    payload = json.loads(Path(file_path).read_text(encoding=encoding))
                    if isinstance(payload, list):
                        df = pd.json_normalize(payload)
                    elif isinstance(payload, dict):
                        df = pd.json_normalize([payload])
                    else:
                        df = pd.DataFrame({'text': [str(payload)]})
            elif suffix == '.parquet':
                df = pd.read_parquet(file_path)
            elif suffix == '.feather':
                df = pd.read_feather(file_path)
            elif suffix in {'.html', '.htm'}:
                tables = pd.read_html(file_path)
                df = tables[0] if tables else pd.DataFrame({'text': [Path(file_path).read_text(errors='replace')]})
            elif suffix == '.pdf':
                from pypdf import PdfReader
                pages = [(page.extract_text() or '') for page in PdfReader(file_path).pages]
                df = pd.DataFrame({'text': [text for text in pages if text.strip()]})
            elif suffix == '.docx':
                from docx import Document
                doc = Document(file_path)
                df = pd.DataFrame({'text': [p.text for p in doc.paragraphs if p.text.strip()]})
            elif suffix == '.pptx':
                from pptx import Presentation
                presentation = Presentation(file_path)
                paragraphs = []
                for slide in presentation.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, 'text') and shape.text.strip():
                            paragraphs.append(shape.text)
                df = pd.DataFrame({'text': paragraphs})
            elif suffix in {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif', '.mp3', '.wav', '.mp4', '.mov', '.avi'}:
                df = pd.DataFrame([{
                    'file_name': os.path.basename(file_path),
                    'file_type': suffix.lstrip('.'),
                    'file_path': str(Path(file_path).resolve()),
                }])
            else:
                return {"success": False, "error": f"不支持的文件格式: {suffix or file_path}"}
            
            self.current_data = df
            self.current_file_path = file_path
            self.current_source_scope = infer_source_scope(file_path, source_scope)
            
            return {
                "success": True,
                "message": f"成功加载 {suffix.upper().lstrip('.')} 文件",
                "file_name": os.path.basename(file_path),
                "rows": len(df),
                "columns": len(df.columns),
                "column_names": df.columns.tolist(),
                "preview": df.head(5).to_dict('records'),
            }
            
        except Exception as e:
            return {"success": False, "error": f"加载数据失败: {str(e)}"}
    
    def preprocess_text(self, 
                       text_column: Optional[str] = None,
                       use_custom_dict: bool = True,
                       use_stopwords: bool = True) -> Dict[str, Any]:
        """文本预处理（分词、去停用词等）"""
        try:
            if self.current_data is None:
                return {"success": False, "error": "请先加载数据"}
                
            if text_column is None:
                text_column = self.auto_detect_text_column()
                if text_column is None:
                    return {"success": False, "error": "无法自动检测到文本列，请明确指定。"}
            
            if text_column not in self.current_data.columns:
                return {"success": False, "error": f"列 '{text_column}' 不存在"}
            
            # Data_PreProcessor 的 __init__ 接受 stopwords_file 和 user_dict_file
            # 为了简化，如果 use_stopwords 为 False，我们可以传递一个不存在的路径或空路径来禁用停用词（具体取决于类的实现，但通常传递None表示使用默认）
            # 为了避免引发错误，我们直接使用默认初始化，或者根据需要传入None。
            # 由于源码中写的是如果为None则使用默认路径，这意味着无法通过传入None来"禁用"停用词。
            # 为了解决当前的 ValueError，我们只传递预期的参数名或不传递。
            
            valid = self.current_data[text_column].notna() & self.current_data[text_column].astype(str).str.strip().ne("")
            texts = self.current_data.loc[valid, text_column].astype(str).tolist()
            self._runtime_log("preprocess_text", f"开始文本预处理；有效文本={len(texts)} 条")
            processed_texts = source_video_text_adapter.preprocess(texts)
            
            new_col_name = f'{text_column}_processed'
            self.current_data[new_col_name] = pd.NA
            self.current_data.loc[self.current_data.index[valid][:len(processed_texts)], new_col_name] = processed_texts

            output_path = self._get_output_path(self.current_file_path, "preprocessed")
            self._runtime_log("preprocess_text", "预处理完成，正在写入结果文件")
            self.current_data.to_csv(output_path, index=False, encoding='utf-8-sig')
            
            return {
                "success": True,
                "message": f"成功对 '{text_column}' 列进行了预处理，生成新列 '{new_col_name}'。",
                "output_file": output_path,
                "preview": self.current_data[[text_column, new_col_name]].head(5).to_dict('records'),
                "processed_columns": [new_col_name],
                "method": "video_text_mutiplemodal_agent.Data_PreProcessor",
                "implementation": "text_processor_subagent_stage_3.common.data_preprocess.Data_PreProcessor",
                "stopwords_removed": True  # 默认为True，因为Data_PreProcessor默认会加载
            }
            
        except Exception as e:
            return {"success": False, "error": f"预处理失败: {str(e)}"}
    
    def sentiment_analysis(self, 
                          text_column: Optional[str] = None,
                          mode: str = "chinese") -> Dict[str, Any]:
        """中文八分类情绪分析，输出源项目兼容的 Excel 结果。"""
        try:
            if self.current_data is None:
                return {"success": False, "error": "请先加载数据"}
                
            if text_column is None:
                text_column = self.auto_detect_text_column()
                if text_column is None:
                    return {"success": False, "error": "无法自动检测到文本列，请明确指定。"}
            
            if text_column not in self.current_data.columns:
                return {"success": False, "error": f"列 '{text_column}' 不存在"}
            
            valid = self.current_data[text_column].notna() & self.current_data[text_column].astype(str).str.strip().ne("")
            texts = self.current_data.loc[valid, text_column].astype(str).tolist()
            if not texts:
                return {"success": False, "error": "没有有效文本数据"}

            if mode == "general":
                self._runtime_log("sentiment_analysis", f"加载通用情感模型；待推理文本={len(texts)} 条")
                results = source_video_text_adapter.general_sentiment(texts)
                result_indexes = self.current_data.index[valid][:len(results)]
                self.current_data["sentiment"] = pd.NA
                self.current_data["sentiment_probability"] = np.nan
                for row_index, prediction in zip(result_indexes, results):
                    self.current_data.at[row_index, "sentiment"] = prediction.get("sentiment", "")
                    self.current_data.at[row_index, "sentiment_probability"] = float(
                        prediction.get("probability", 0.0)
                    )
                source_compatible = pd.DataFrame(results, columns=['text', 'sentiment', 'probability'])
                summary = (source_compatible.groupby('sentiment', dropna=False)
                           .agg(count=('text', 'size'), mean_probability=('probability', 'mean'))
                           .reset_index().sort_values('count', ascending=False))
                summary['percentage'] = summary['count'] / max(len(source_compatible), 1)
                run_dir = create_run_dir(self.current_file_path, "general_sentiment", self.current_source_scope)
                output_path = run_dir / f"{Path(self.current_file_path).stem}_general_sentiment.xlsx"
                self._runtime_log("sentiment_analysis", "模型推理完成，正在生成情感分析工作簿")
                write_workbook(output_path, {'情感分析结果': source_compatible, '情感分布': summary})
                return {
                    "success": True,
                    "message": "通用五分类情感分析完成。",
                    "output_file": str(output_path),
                    "artifacts": artifact_result([output_path]),
                    "preview": source_compatible.head(5).to_dict('records'),
                    "sentiment_distribution": source_compatible['sentiment'].value_counts().to_dict(),
                    "model": "video_text_mutiplemodal_agent/sentiment_general_analysis",
                    "implementation": "academic_agent.integrations.video_text_adapter.SourceVideoTextAdapter.general_sentiment",
                }
            if mode != "chinese":
                return {"success": False, "error": "mode 仅支持 chinese 或 general"}
            
            self._runtime_log("sentiment_analysis", f"加载中文情绪模型；待推理文本={len(texts)} 条；batch_size=8")
            results = source_video_text_adapter.sentiment(texts, top_k=1, batch_size=8)
            
            result_indexes = self.current_data.index[valid][:len(results)]
            self.current_data['emotion_en'] = pd.NA
            self.current_data['emotion_zh'] = pd.NA
            self.current_data['probability'] = np.nan
            for row_index, prediction in zip(result_indexes, results):
                self.current_data.at[row_index, 'emotion_en'] = prediction.get('primary_emotion_en', 'none')
                self.current_data.at[row_index, 'emotion_zh'] = prediction.get('primary_emotion_zh', '无情绪')
                self.current_data.at[row_index, 'probability'] = float(prediction.get('primary_probability', 0.0))

            id_column = self._detect_id_column()
            ids = self.current_data.loc[result_indexes, id_column] if id_column else pd.Series(result_indexes, index=result_indexes)
            source_compatible = pd.DataFrame({
                'cid': ids.values,
                'text': self.current_data.loc[result_indexes, text_column].values,
                'emotion_en': self.current_data.loc[result_indexes, 'emotion_en'].values,
                'emotion_zh': self.current_data.loc[result_indexes, 'emotion_zh'].values,
                'probability': self.current_data.loc[result_indexes, 'probability'].values,
            })
            summary = (source_compatible.groupby(['emotion_en', 'emotion_zh'], dropna=False)
                       .agg(count=('text', 'size'), mean_probability=('probability', 'mean'))
                       .reset_index().sort_values('count', ascending=False))
            summary['percentage'] = summary['count'] / max(len(source_compatible), 1)
            run_dir = create_run_dir(self.current_file_path, "chinese_emotion", self.current_source_scope)
            output_path = run_dir / f"{Path(self.current_file_path).stem}_chinese_emotion.xlsx"
            self._runtime_log("sentiment_analysis", "模型推理完成，正在生成情绪结果和统计工作簿")
            write_workbook(output_path, {
                '情感分析结果': source_compatible,
                '情感分布': summary,
                '原始数据与结果': self.current_data.copy(),
            })
            emotion_counts = source_compatible['emotion_zh'].value_counts().to_dict()

            return {
                "success": True,
                "message": "情感分析完成。",
                "output_file": str(output_path),
                "artifacts": artifact_result([output_path]),
                "preview": source_compatible.head(5).to_dict('records'),
                "sentiment_distribution": emotion_counts,
                "model": "video_text_mutiplemodal_agent/chinese_sentiment_bert_analysis",
                "implementation": "text_processor_subagent_stage_3.sentiment.chinese_sentiment_bert_analysis.predict_emotion",
            }
            
        except Exception as e:
            return {"success": False, "error": f"情感分析失败: {str(e)}"}
    
    def text_clustering(self,
                       text_column: Optional[str] = None,
                       n_clusters: int = 5,
                       algorithm: str = "kmeans",
                       eps: float = 0.5,
                       min_samples: int = 5) -> Dict[str, Any]:
        """BGE 文本聚类，保留源项目 cid/text/cluster_id 主结果结构。"""
        try:
            if self.current_data is None:
                return {"success": False, "error": "请先加载数据"}
                
            if text_column is None:
                text_column = self.auto_detect_text_column()
                if text_column is None:
                    return {"success": False, "error": "无法自动检测到文本列，请明确指定。"}
            
            if text_column not in self.current_data.columns:
                return {"success": False, "error": f"列 '{text_column}' 不存在"}
            
            texts = self.current_data[text_column].dropna().astype(str).tolist()
            if not texts:
                return {"success": False, "error": "没有有效文本数据"}

            self._runtime_log(
                "text_clustering",
                f"加载 BGE 与 {algorithm} 聚类算法；文本={len(texts)} 条；聚类数={n_clusters}",
            )
            clustered = source_video_text_adapter.clustering(
                self.current_data,
                text_column=text_column,
                n_clusters=n_clusters,
                algorithm=algorithm,
                eps=eps,
                min_samples=min_samples,
            )
            self.current_data["cluster_id"] = pd.NA
            for _, row in clustered.iterrows():
                self.current_data.iloc[int(row["row_id"]), self.current_data.columns.get_loc("cluster_id")] = int(row["cluster"])
            id_column = self._detect_id_column()
            valid = self.current_data['cluster_id'].notna()
            ids = self.current_data.loc[valid, id_column] if id_column else pd.Series(self.current_data.index[valid], index=self.current_data.index[valid])
            source_compatible = pd.DataFrame({
                'cid': ids.values,
                'text': self.current_data.loc[valid, text_column].values,
                'cluster_id': self.current_data.loc[valid, 'cluster_id'].astype(int).values,
            })
            summary = (source_compatible.groupby('cluster_id')
                       .agg(segment_count=('text', 'size'), unique_text_count=('text', 'nunique'))
                       .reset_index())
            summary['percentage'] = summary['segment_count'] / max(len(source_compatible), 1)
            run_dir = create_run_dir(self.current_file_path, "clustering", self.current_source_scope)
            csv_path = run_dir / f"{Path(self.current_file_path).stem}_clustering_results.csv"
            excel_path = run_dir / f"{Path(self.current_file_path).stem}_clustering_results.xlsx"
            self._runtime_log("text_clustering", "向量编码与聚类完成，正在写入 CSV 和 Excel")
            source_compatible.to_csv(csv_path, index=False, encoding='utf-8-sig')
            write_workbook(excel_path, {
                '聚类结果': source_compatible,
                '聚类统计': summary,
                '原始数据与结果': self.current_data.copy(),
            })
            cluster_counts = source_compatible['cluster_id'].value_counts().to_dict()
            cluster_distribution = []
            cluster_profiles = []
            for cluster_id, group in source_compatible.groupby('cluster_id', sort=True):
                cluster_id = int(cluster_id)
                segment_count = int(len(group))
                percentage = round(segment_count / max(len(source_compatible), 1), 6)
                representative_texts = [
                    str(value).replace("\n", " ").strip()[:180]
                    for value in group['text'].head(3).tolist()
                    if str(value).strip()
                ]
                cluster_distribution.append({
                    "cluster_id": cluster_id,
                    "segment_count": segment_count,
                    "percentage": percentage,
                })
                cluster_profiles.append({
                    "cluster_id": cluster_id,
                    "segment_count": segment_count,
                    "percentage": percentage,
                    "representative_texts": representative_texts,
                })
            distribution_text = "；".join(
                f"簇{item['cluster_id']}={item['segment_count']}条（{item['percentage']:.2%}）"
                for item in cluster_distribution
            )

            return {
                "success": True,
                "message": f"文本聚类完成，使用 {algorithm}，共得到 {source_compatible['cluster_id'].nunique()} 个簇。",
                "output_file": str(excel_path),
                "output_files": artifact_result([excel_path, csv_path]),
                "artifacts": artifact_result([excel_path, csv_path]),
                "preview": source_compatible.head(5).to_dict('records'),
                "algorithm": f"BGE/SentenceTransformer + {algorithm}",
                "implementation": "text_processor_subagent_stage_3.common.textEmbedding.TextEmbedding + sklearn.cluster",
                "algorithm_name": str(algorithm).lower(),
                "n_clusters": n_clusters,
                "cluster_sizes": [int(cluster_counts.get(i, 0)) for i in sorted(cluster_counts)],
                # 这些结构化摘要随工具观察返回给 Agent，避免主模型只能看到
                # 输出文件名和前几行样例，进而生成被截断或不完整的聚类说明。
                "cluster_distribution": cluster_distribution,
                "cluster_profiles": cluster_profiles,
                "summary_for_agent": (
                    f"共得到 {len(cluster_distribution)} 个主题簇；完整分布：{distribution_text}。"
                    "每个簇的代表文本已在 cluster_profiles 中提供，可据此概括主题特征。"
                ),
            }
            
        except Exception as e:
            return {"success": False, "error": f"聚类失败: {str(e)}"}

    def repair_text_clustering(self, text_column: Optional[str] = None) -> Dict[str, Any]:
        """按已有聚类标签重建结果，不重新生成文本向量，也不调用 BERT/BGE。"""
        try:
            if self.current_data is None:
                return {"success": False, "error": "请先加载数据", "repair_mode": True, "used_model": False}

            label_column = next(
                (
                    column for column in ("cluster_id", "cluster", "New Cluster")
                    if column in self.current_data.columns
                ),
                None,
            )
            if label_column is None:
                return {
                    "success": False,
                    "error": "当前数据没有已有聚类标签，规则修复无法重建结果；请重新发起一次聚类分析。",
                    "repair_mode": True,
                    "used_model": False,
                }
            text_column = text_column or self.auto_detect_text_column()
            if not text_column or text_column not in self.current_data.columns:
                return {
                    "success": False,
                    "error": "规则修复找不到原始文本列，无法重建聚类特征描述。",
                    "repair_mode": True,
                    "used_model": False,
                }

            working = self.current_data[[text_column, label_column]].copy()
            working = working.dropna(subset=[text_column, label_column])
            working[text_column] = working[text_column].astype(str).str.strip()
            working = working[working[text_column].ne("")]
            if working.empty:
                return {
                    "success": False,
                    "error": "已有聚类标签下没有可用文本，规则修复无法继续。",
                    "repair_mode": True,
                    "used_model": False,
                }

            id_column = self._detect_id_column()
            ids = (
                self.current_data.loc[working.index, id_column]
                if id_column else pd.Series(working.index, index=working.index)
            )
            source_compatible = pd.DataFrame({
                "cid": ids.values,
                "text": working[text_column].values,
                "cluster_id": working[label_column].values,
            })
            summary = (
                source_compatible.groupby("cluster_id", dropna=False)
                .agg(segment_count=("text", "size"), unique_text_count=("text", "nunique"))
                .reset_index()
            )
            summary["percentage"] = summary["segment_count"] / max(len(source_compatible), 1)

            cluster_distribution: list[dict[str, Any]] = []
            cluster_profiles: list[dict[str, Any]] = []
            for cluster_id, group in source_compatible.groupby("cluster_id", sort=True, dropna=False):
                representative_texts = [
                    str(value).replace("\n", " ").strip()[:180]
                    for value in group["text"].head(3).tolist()
                    if str(value).strip()
                ]
                count = int(len(group))
                cluster_distribution.append({
                    "cluster_id": cluster_id,
                    "segment_count": count,
                    "percentage": round(count / max(len(source_compatible), 1), 6),
                })
                cluster_profiles.append({
                    "cluster_id": cluster_id,
                    "segment_count": count,
                    "percentage": round(count / max(len(source_compatible), 1), 6),
                    "representative_texts": representative_texts,
                    "feature_description": "；".join(representative_texts[:2]) or "暂无代表文本",
                })

            run_dir = create_run_dir(self.current_file_path, "clustering_repair", self.current_source_scope)
            csv_path = run_dir / f"{Path(self.current_file_path).stem}_clustering_repaired.csv"
            excel_path = run_dir / f"{Path(self.current_file_path).stem}_clustering_repaired.xlsx"
            source_compatible.to_csv(csv_path, index=False, encoding="utf-8-sig")
            write_workbook(excel_path, {
                "聚类修复结果": source_compatible,
                "聚类分布": summary,
                "原始数据与结果": self.current_data.copy(),
            })
            distribution_text = "；".join(
                f"簇{item['cluster_id']}={item['segment_count']}条（{item['percentage']:.2%}）"
                for item in cluster_distribution
            )
            return {
                "success": True,
                "message": "已按已有聚类标签重建分布和特征描述。",
                "output_file": str(excel_path),
                "output_files": artifact_result([excel_path, csv_path]),
                "artifacts": artifact_result([excel_path, csv_path]),
                "algorithm": "rule_based_repair",
                "implementation": "academic_agent.tools.text_mining_tools.TextMiningTools.repair_text_clustering",
                "repair_mode": "cluster_structure_repair",
                "used_model": False,
                "cluster_distribution": cluster_distribution,
                "cluster_profiles": cluster_profiles,
                "n_clusters": len(cluster_distribution),
                "summary_for_agent": f"规则修复完成；已有标签下共 {len(cluster_distribution)} 个主题簇，完整分布：{distribution_text}。",
            }
        except Exception as exc:
            return {
                "success": False,
                "error": f"聚类规则修复失败：{exc}",
                "repair_mode": True,
                "used_model": False,
            }

    def repair_sentiment_analysis(self, text_column: Optional[str] = None) -> Dict[str, Any]:
        """按已有情感标签重建结果，不重新推理，也不调用任何 BERT 模型。"""
        try:
            if self.current_data is None:
                return {"success": False, "error": "请先加载数据", "repair_mode": True, "used_model": False}

            label_column = next(
                (
                    column for column in ("emotion_zh", "sentiment", "predicted_sentiment", "emotion_en")
                    if column in self.current_data.columns
                    and self.current_data[column].notna().astype(bool).any()
                ),
                None,
            )
            if label_column is None:
                return {
                    "success": False,
                    "error": "当前数据没有已有情感标签，规则修复无法补造分类结果；请重新发起一次情感分析。",
                    "repair_mode": True,
                    "used_model": False,
                }
            text_column = text_column or self.auto_detect_text_column()
            if not text_column or text_column not in self.current_data.columns:
                return {
                    "success": False,
                    "error": "规则修复找不到原始文本列，无法重建情感结果。",
                    "repair_mode": True,
                    "used_model": False,
                }

            working = self.current_data[[text_column, label_column]].copy()
            working = working.dropna(subset=[text_column, label_column])
            working[text_column] = working[text_column].astype(str).str.strip()
            working[label_column] = working[label_column].astype(str).str.strip()
            working = working[working[text_column].ne("") & working[label_column].ne("")]
            if working.empty:
                return {
                    "success": False,
                    "error": "已有情感标签下没有可用文本，规则修复无法继续。",
                    "repair_mode": True,
                    "used_model": False,
                }

            id_column = self._detect_id_column()
            ids = (
                self.current_data.loc[working.index, id_column]
                if id_column else pd.Series(working.index, index=working.index)
            )
            probability_column = next(
                (
                    column for column in ("probability", "sentiment_probability")
                    if column in self.current_data.columns
                ),
                None,
            )
            result_data = {
                "cid": ids.values,
                "text": working[text_column].values,
                "sentiment": working[label_column].values,
            }
            if probability_column:
                result_data["probability"] = pd.to_numeric(
                    self.current_data.loc[working.index, probability_column], errors="coerce"
                ).values
            source_compatible = pd.DataFrame(result_data)
            aggregation = {"count": ("text", "size")}
            if "probability" in source_compatible:
                aggregation["mean_probability"] = ("probability", "mean")
            summary = (
                source_compatible.groupby("sentiment", dropna=False)
                .agg(**aggregation)
                .reset_index()
                .sort_values("count", ascending=False)
            )
            summary["percentage"] = summary["count"] / max(len(source_compatible), 1)

            run_dir = create_run_dir(self.current_file_path, "sentiment_repair", self.current_source_scope)
            output_path = run_dir / f"{Path(self.current_file_path).stem}_sentiment_repaired.xlsx"
            write_workbook(excel_path := output_path, {
                "情感修复结果": source_compatible,
                "情感分布": summary,
                "原始数据与结果": self.current_data.copy(),
            })
            return {
                "success": True,
                "message": "已按已有情感标签重建情感分布和结果工作簿。",
                "output_file": str(excel_path),
                "artifacts": artifact_result([excel_path]),
                "sentiment_distribution": source_compatible["sentiment"].value_counts().to_dict(),
                "implementation": "academic_agent.tools.text_mining_tools.TextMiningTools.repair_sentiment_analysis",
                "repair_mode": "sentiment_structure_repair",
                "used_model": False,
                "summary_for_agent": "规则修复完成；仅复用已有情感标签和概率重建结果，没有重新调用 BERT。",
            }
        except Exception as exc:
            return {
                "success": False,
                "error": f"情感规则修复失败：{exc}",
                "repair_mode": True,
                "used_model": False,
            }
    
    def extract_keywords(self,
                        text_column: Optional[str] = None,
                        top_n: int = 10) -> Dict[str, Any]:
        """使用源项目 KeyBERT 融合算法按聚类或全文提取关键词并生成工作簿。"""
        try:
            if self.current_data is None:
                return {"success": False, "error": "请先加载数据"}
                
            if text_column is None:
                text_column = self.auto_detect_text_column()
                if text_column is None:
                    return {"success": False, "error": "无法自动检测到文本列，请明确指定。"}
            
            if text_column not in self.current_data.columns:
                return {"success": False, "error": f"列 '{text_column}' 不存在"}
            
            group_column = next((column for column in ('New Cluster', 'cluster_id', 'cluster') if column in self.current_data.columns), None)
            id_column = self._detect_id_column()
            method = "video_text_mutiplemodal_agent.TextKeyBert"
            self._runtime_log(
                "extract_keywords",
                f"加载 KeyBERT 关键词算法；分组字段={group_column or '全文'}；TopN={top_n}",
            )
            keyword_rows = source_video_text_adapter.keywords_by_group(
                self.current_data, text_column, group_column, id_column, top_n
            )
            for row in keyword_rows:
                row['source_file'] = Path(self.current_file_path).name
            keyword_table = pd.DataFrame(keyword_rows, columns=['source_file', 'cluster', 'segment_count', 'comment_count', 'keywords'])
            long_rows = []
            for row in keyword_rows:
                try:
                    pairs = ast.literal_eval(row['keywords'])
                except (ValueError, SyntaxError):
                    pairs = []
                for rank, pair in enumerate(pairs, 1):
                    if isinstance(pair, (list, tuple)) and len(pair) >= 2:
                        long_rows.append({'cluster': row['cluster'], 'rank': rank, 'keyword': pair[0], 'score': pair[1]})
            long_table = pd.DataFrame(long_rows, columns=['cluster', 'rank', 'keyword', 'score'])
            run_dir = create_run_dir(self.current_file_path, "keywords", self.current_source_scope)
            output_path = run_dir / f"{Path(self.current_file_path).stem}_keywords.xlsx"
            self._runtime_log("extract_keywords", "关键词计算完成，正在生成结果工作簿")
            write_workbook(output_path, {'keywords': keyword_table, '关键词明细': long_table})
            return {
                "success": True,
                "message": f"从 '{text_column}' 列中提取了 Top {top_n} 关键词。",
                "method": method,
                "implementation": "text_processor_subagent_stage_3.clustering.textKeyBert_cluster_sentence_fuse_agent.TextKeyBert",
                "custom_dictionary": str(source_video_text_adapter.custom_dictionary_path) if source_video_text_adapter.custom_dictionary_path else None,
                "output_file": str(output_path),
                "artifacts": artifact_result([output_path]),
                "keywords": long_table.head(top_n).to_dict('records'),
            }
            
        except Exception as e:
            return {"success": False, "error": f"关键词提取失败: {str(e)}"}

    def _llm_information_extraction(
        self,
        text_column: Optional[str],
        extraction_type: str,
        batch_size: int = 8,
        max_texts: int = 200,
        provider: Optional[str] = None,
        model: Optional[str] = None,
        base_url: Optional[str] = None,
        api_key: Optional[str] = None,
    ) -> Dict[str, Any]:
        """按行调用 OpenAI-compatible 大模型并生成实体/关系工作簿。"""
        try:
            if self.current_data is None:
                return {"success": False, "error": "请先加载数据"}
            text_column = text_column or self.auto_detect_text_column()
            if not text_column or text_column not in self.current_data.columns:
                return {"success": False, "error": "无法找到文本列，请明确指定 text_column。"}

            valid = (
                self.current_data[text_column].notna()
                & self.current_data[text_column].astype(str).str.strip().ne("")
            )
            indexes = list(self.current_data.index[valid])[: max(1, int(max_texts))]
            texts = [str(self.current_data.at[index, text_column]).strip() for index in indexes]
            if not texts:
                return {"success": False, "error": "没有有效文本数据"}

            from academic_agent.algorithms.text_mining.information_extraction import LLMInformationExtractor

            self._runtime_log(
                "llm_information_extraction",
                f"调用大模型进行信息抽取；类型={extraction_type}；文本={len(texts)} 条；batch_size={batch_size}",
            )
            extractor = LLMInformationExtractor(
                provider=provider, model=model, base_url=base_url, api_key=api_key
            )
            batch = extractor.extract_batch(texts, extraction_type=extraction_type, batch_size=batch_size)

            id_column = self._detect_id_column()
            record_ids = [self.current_data.at[index, id_column] if id_column else index for index in indexes]
            entity_rows: list[dict[str, Any]] = []
            relation_rows: list[dict[str, Any]] = []
            entity_json_by_index: dict[Any, str] = {}
            relation_json_by_index: dict[Any, str] = {}
            for original_index, record_id, item in zip(indexes, record_ids, batch.items):
                entities = item.get("entities", [])
                relations = item.get("relations", [])
                selected_entities = entities if extraction_type in {"entities", "both"} else []
                selected_relations = relations if extraction_type in {"relations", "both"} else []
                entity_json_by_index[original_index] = json.dumps(selected_entities, ensure_ascii=False)
                relation_json_by_index[original_index] = json.dumps(selected_relations, ensure_ascii=False)
                for entity in selected_entities:
                    entity_rows.append({
                        "record_id": record_id,
                        "row_index": original_index,
                        "text": item["text"],
                        "entity_text": entity["text"],
                        "entity_type": entity["type"],
                        "start": entity.get("start"),
                        "end": entity.get("end"),
                        "confidence": entity.get("confidence"),
                        "llm_model": batch.model,
                    })
                for relation in selected_relations:
                    relation_rows.append({
                        "record_id": record_id,
                        "row_index": original_index,
                        "text": item["text"],
                        "subject": relation["subject"],
                        "subject_type": relation["subject_type"],
                        "predicate": relation["predicate"],
                        "object": relation["object"],
                        "object_type": relation["object_type"],
                        "confidence": relation.get("confidence"),
                        "llm_model": batch.model,
                    })

            if extraction_type in {"entities", "both"}:
                self.current_data["entities_json"] = pd.NA
                for index, value in entity_json_by_index.items():
                    self.current_data.at[index, "entities_json"] = value
            if extraction_type in {"relations", "both"}:
                self.current_data["relations_json"] = pd.NA
                for index, value in relation_json_by_index.items():
                    self.current_data.at[index, "relations_json"] = value

            entity_table = pd.DataFrame(entity_rows, columns=[
                "record_id", "row_index", "text", "entity_text", "entity_type",
                "start", "end", "confidence", "llm_model",
            ])
            relation_table = pd.DataFrame(relation_rows, columns=[
                "record_id", "row_index", "text", "subject", "subject_type", "predicate",
                "object", "object_type", "confidence", "llm_model",
            ])
            task_name = "entity_recognition" if extraction_type == "entities" else "relation_extraction"
            if extraction_type == "both":
                task_name = "information_extraction"
            run_dir = create_run_dir(self.current_file_path, task_name, self.current_source_scope)
            output_path = run_dir / f"{Path(self.current_file_path or 'analysis').stem}_{task_name}.xlsx"
            sheets = {"原始数据与结果": self.current_data.copy()}
            if extraction_type in {"entities", "both"}:
                sheets["实体识别结果"] = entity_table
            if extraction_type in {"relations", "both"}:
                sheets["关系抽取结果"] = relation_table
            write_workbook(output_path, sheets)
            self._runtime_log(
                "llm_information_extraction",
                f"信息抽取完成；实体={len(entity_table)} 条；关系={len(relation_table)} 条",
            )
            return {
                "success": True,
                "message": f"大模型{('实体识别' if extraction_type == 'entities' else '关系抽取' if extraction_type == 'relations' else '实体识别和关系抽取')}完成。",
                "output_file": str(output_path),
                "artifacts": artifact_result([output_path]),
                "text_column": text_column,
                "processed_texts": len(texts),
                "entity_count": len(entity_table),
                "relation_count": len(relation_table),
                "entity_types": entity_table["entity_type"].value_counts().to_dict() if not entity_table.empty else {},
                "relation_types": relation_table["predicate"].value_counts().to_dict() if not relation_table.empty else {},
                "entities": entity_table.head(20).to_dict("records"),
                "relations": relation_table.head(20).to_dict("records"),
                "model": batch.model,
                "provider": batch.provider,
                "llm_used": True,
                "implementation": "academic_agent.algorithms.text_mining.information_extraction.LLMInformationExtractor",
            }
        except Exception as exc:
            return {"success": False, "error": f"大模型信息抽取失败: {type(exc).__name__}: {exc}", "llm_used": True}

    def entity_recognition(
        self,
        text_column: Optional[str] = None,
        batch_size: int = 8,
        max_texts: int = 200,
        provider: Optional[str] = None,
        model: Optional[str] = None,
        base_url: Optional[str] = None,
        api_key: Optional[str] = None,
    ) -> Dict[str, Any]:
        """使用大模型识别实体，结果同时写入实体明细和当前数据。"""
        return self._llm_information_extraction(
            text_column, "entities", batch_size, max_texts, provider, model, base_url, api_key
        )

    def relation_extraction(
        self,
        text_column: Optional[str] = None,
        batch_size: int = 8,
        max_texts: int = 200,
        provider: Optional[str] = None,
        model: Optional[str] = None,
        base_url: Optional[str] = None,
        api_key: Optional[str] = None,
    ) -> Dict[str, Any]:
        """使用大模型抽取主语-关系-宾语三元组。"""
        return self._llm_information_extraction(
            text_column, "relations", batch_size, max_texts, provider, model, base_url, api_key
        )

# 创建全局工具实例
text_mining_tools = TextMiningTools()
