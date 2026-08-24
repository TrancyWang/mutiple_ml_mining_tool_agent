"""多格式学术文档生成器。

输入统一使用 Agent 生成的 Markdown/纯文本内容，根据目标扩展名生成
对应的文本、Word、PDF、PPT 或 Excel 文件。生成结果以内存 bytes 返回，
由 WorkspaceManager 统一负责路径安全、权限和确认流程。
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any


TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".json", ".jsonl", ".csv", ".tsv",
    ".html", ".htm", ".xml", ".yaml", ".yml", ".log", ".tex",
}
BINARY_EXTENSIONS = {".docx", ".pdf", ".pptx", ".xlsx"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | BINARY_EXTENSIONS


def _content_bytes(content: str | bytes) -> bytes:
    return content if isinstance(content, bytes) else str(content).encode("utf-8")


def _markdown_blocks(content: str) -> list[tuple[str, str]]:
    """将简单 Markdown 转换成文档生成器共用的块结构。"""
    blocks: list[tuple[str, str]] = []
    paragraph: list[str] = []
    for raw_line in str(content).splitlines():
        line = raw_line.strip()
        if not line:
            if paragraph:
                blocks.append(("paragraph", " ".join(paragraph)))
                paragraph = []
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            if paragraph:
                blocks.append(("paragraph", " ".join(paragraph)))
                paragraph = []
            blocks.append((f"heading{len(heading.group(1))}", heading.group(2).strip()))
        elif re.match(r"^[-*+]\s+", line):
            if paragraph:
                blocks.append(("paragraph", " ".join(paragraph)))
                paragraph = []
            blocks.append(("bullet", re.sub(r"^[-*+]\s+", "", line)))
        elif re.match(r"^\d+[.)]\s+", line):
            if paragraph:
                blocks.append(("paragraph", " ".join(paragraph)))
                paragraph = []
            blocks.append(("number", re.sub(r"^\d+[.)]\s+", "", line)))
        else:
            paragraph.append(line)
    if paragraph:
        blocks.append(("paragraph", " ".join(paragraph)))
    return blocks or [("paragraph", "")]


def _generate_docx(content: str) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("生成 Word 文档需要安装 python-docx") from exc

    document = Document()
    for kind, text in _markdown_blocks(content):
        if kind.startswith("heading"):
            level = min(int(kind[-1]), 6)
            document.add_heading(text, level=level)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        elif kind == "number":
            document.add_paragraph(text, style="List Number")
        else:
            document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _generate_pptx(content: str) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("生成 PPT 文档需要安装 python-pptx") from exc

    blocks = _markdown_blocks(content)
    presentation = Presentation()
    presentation.core_properties.title = "Academic Agent 生成演示文稿"
    presentation.core_properties.subject = "Academic Agent"
    chunks: list[list[tuple[str, str]]] = []
    current: list[tuple[str, str]] = []
    for block in blocks:
        current.append(block)
        if len(current) >= 12:
            chunks.append(current)
            current = []
    if current or not chunks:
        chunks.append(current)

    for index, chunk in enumerate(chunks):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        title = next((text for kind, text in chunk if kind.startswith("heading")), None)
        slide.shapes.title.text = title or f"学术分析结果（第 {index + 1} 页）"
        title_frame = slide.shapes.title.text_frame
        for paragraph in title_frame.paragraphs:
            paragraph.font.size = Pt(26)
        body = slide.placeholders[1].text_frame
        body.clear()
        for kind, text in chunk:
            if title and text == title and kind.startswith("heading"):
                continue
            paragraph = body.paragraphs[0] if len(body.paragraphs) == 1 and not body.paragraphs[0].text else body.add_paragraph()
            paragraph.text = ("• " if kind == "bullet" else "1. " if kind == "number" else "") + text
            paragraph.level = 0
            paragraph.font.size = Pt(18)
        body.margin_left = Inches(0.25)

    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _generate_pdf(content: str) -> bytes:
    """使用项目已有的 Matplotlib 后端生成分页 PDF，避免强制新增 PDF 库。"""
    try:
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        from matplotlib.backends.backend_pdf import PdfPages
        from matplotlib.font_manager import FontProperties
        from matplotlib.figure import Figure
    except ImportError as exc:
        raise RuntimeError("生成 PDF 文档需要安装 matplotlib") from exc

    lines: list[str] = []
    for raw_line in str(content).splitlines() or [""]:
        if not raw_line.strip():
            lines.append("")
            continue
        lines.extend(
            [raw_line[i:i + 58] for i in range(0, len(raw_line), 58)] or [""]
        )
    lines_per_page = 42
    font_candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    font_path = next((path for path in font_candidates if Path(path).is_file()), None)
    font_properties = FontProperties(fname=font_path) if font_path else None
    output = io.BytesIO()
    with PdfPages(output) as pdf:
        for offset in range(0, max(len(lines), 1), lines_per_page):
            page_lines = lines[offset:offset + lines_per_page]
            figure = Figure(figsize=(8.27, 11.69), dpi=150)
            FigureCanvasAgg(figure)
            figure.text(
                0.08,
                0.95,
                "\n".join(page_lines),
                va="top",
                ha="left",
                fontsize=10,
                linespacing=1.45,
                fontproperties=font_properties,
            )
            figure.text(
                0.08,
                0.035,
                f"Academic Agent  |  第 {offset // lines_per_page + 1} 页",
                fontsize=8,
                fontproperties=font_properties,
            )
            pdf.savefig(figure, bbox_inches="tight")
    return output.getvalue()


def _generate_xlsx(content: str) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("生成 Excel 文档需要安装 openpyxl") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Academic Agent"
    rows: list[list[Any]] = []
    try:
        parsed = json.loads(content)
        if isinstance(parsed, list) and all(isinstance(row, list) for row in parsed):
            rows = parsed
    except (TypeError, json.JSONDecodeError):
        pass
    if not rows:
        delimiter = "\t" if "\t" in content else "," if "," in content else None
        rows = list(csv.reader(io.StringIO(content), delimiter=delimiter or ",")) if delimiter else [[line] for line in content.splitlines()]
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_document(path: str | Path, content: str | bytes) -> bytes:
    """根据路径扩展名生成文件 bytes。"""
    suffix = Path(path).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return _content_bytes(content)
    if suffix == ".docx":
        return _generate_docx(str(content))
    if suffix == ".pptx":
        return _generate_pptx(str(content))
    if suffix == ".pdf":
        return _generate_pdf(str(content))
    if suffix == ".xlsx":
        return _generate_xlsx(str(content))
    supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
    raise ValueError(f"暂不支持生成 {suffix or '无扩展名'} 文件。支持格式：{supported}")
