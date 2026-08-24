"""Consistent, workspace-scoped outputs for analysis tools."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Mapping
import unicodedata

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from academic_agent.infrastructure.runtime_paths import output_root
from academic_agent.infrastructure.workspace_manager import workspace_manager


GREEN = "18794E"
LIGHT_GREEN = "EAF6EF"
GRID = "D8E5DD"


def _display_width(value: object) -> int:
    text = "" if value is None else str(value)
    return sum(2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1 for character in text)


def infer_source_scope(source_file: str | None, source_scope: str | None = None) -> str:
    """Return the output scope without confusing uploaded files with workspace files."""
    if source_scope in {"upload", "workspace"}:
        return source_scope
    if not source_file:
        return "upload"
    candidate = Path(source_file).expanduser()
    if not candidate.is_absolute():
        candidate = workspace_manager.root / candidate
    try:
        return "workspace" if candidate.resolve().is_relative_to(workspace_manager.root) else "upload"
    except (OSError, RuntimeError):
        return "upload"


def artifact_output_root(source_file: str | None = None, source_scope: str | None = None) -> Path:
    """Resolve the product output policy for workspace files and uploaded files."""
    if infer_source_scope(source_file, source_scope) == "workspace":
        root = workspace_manager.root / "output"
    else:
        root = output_root()
    root.mkdir(parents=True, exist_ok=True)
    return root.resolve()


def create_run_dir(source_file: str | None, task: str, source_scope: str | None = None) -> Path:
    stem = Path(source_file or "analysis").stem
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:19]
    path = artifact_output_root(source_file, source_scope) / f"{stem}_{task}_{timestamp}"
    path.mkdir(parents=True, exist_ok=False)
    return path


def write_workbook(path: str | Path, sheets: Mapping[str, pd.DataFrame]) -> Path:
    """Write machine-readable tables with restrained academic formatting."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, dataframe in sheets.items():
            dataframe.to_excel(writer, sheet_name=str(sheet_name)[:31], index=False)
        workbook = writer.book
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = False
            header_fill = PatternFill("solid", fgColor=GREEN)
            header_font = Font(color="FFFFFF", bold=True)
            thin = Side(style="thin", color=GRID)
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")
                cell.border = Border(bottom=thin)
            worksheet.row_dimensions[1].height = 24
            for column_index in range(1, worksheet.max_column + 1):
                values = [worksheet.cell(row, column_index).value for row in range(1, min(worksheet.max_row, 300) + 1)]
                width = min(60, max(10, max((_display_width(value) for value in values if value is not None), default=8) + 2))
                worksheet.column_dimensions[get_column_letter(column_index)].width = width
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=False)
    return path


def artifact_result(paths: list[str | Path]) -> list[str]:
    return [str(Path(path).resolve()) for path in paths if Path(path).exists()]
